"""Structural and content coverage for publication-ready Word results reports."""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pandas as pd
import pytest

from biostat_service.analyses import run_plan
from biostat_service.contracts import AnalysisPlan, PlanItem, StudyBrief
from biostat_service.projects import LocalProject
from biostat_service.reporting import build_results_docx, format_p_value
from biostat_service.visuals import build_figures


FIXTURE_PATH = Path(__file__).resolve().parents[3] / "tests" / "fixtures" / "core-study.xlsx"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def _complete_docx_payload(path: Path) -> bytes:
    """Return every uncompressed ZIP member plus its name for leak detection."""
    with ZipFile(path) as archive:
        return b"\n".join(
            name.encode("utf-8") + b"\n" + archive.read(name)
            for name in sorted(archive.namelist())
        )


def _document_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(node.text or "" for node in document.element.iter(qn("w:t")))


def _item(identifier: str, method: str, variables: list[str]) -> PlanItem:
    return PlanItem(
        id=identifier,
        estimand="Confirmed structured estimand.",
        method=method,
        rationale="Confirmed structured rationale.",
        required_variables=variables,
        assumptions=["Confirmed assumptions."],
        robust_alternative="mann_whitney_u" if method == "welch_t_test" else None,
        outputs=["effect_size:required", "confidence_interval:95_percent"],
    )


@pytest.fixture
def report_inputs(tmp_path: Path):
    frame = pd.read_excel(FIXTURE_PATH, sheet_name="Analysis")
    plan = AnalysisPlan(
        version=3,
        items=[
            _item(
                "descriptive_summary",
                "descriptive_summary",
                ["age_years", "treatment_group"],
            ),
            _item(
                "primary_outcome",
                "welch_t_test",
                ["age_years", "treatment_group"],
            ),
        ],
        warnings=["review_assumptions"],
    )
    brief = StudyBrief(
        title="Cardiovascular outcomes",
        question="Is treatment associated with age at the measured endpoint?",
        hypothesis="The exposure groups differ at the measured endpoint.",
        design="cohort",
        outcome_variables=["age_years"],
        exposure_variables=["treatment_group"],
        language="en",
    )
    bundle = run_plan(frame, plan)
    figures = build_figures(frame, plan, bundle, tmp_path / "figures", "en")
    return LocalProject(tmp_path / "cardio.biostat"), brief, plan, bundle, figures


@pytest.fixture
def report_en(tmp_path: Path, report_inputs) -> Path:
    project, brief, plan, bundle, figures = report_inputs
    return build_results_docx(
        project,
        brief,
        plan,
        bundle,
        figures,
        "en",
        tmp_path / "results-en.docx",
    )


def _style_properties(document: Document, style_name: str) -> dict[str, str | None]:
    style = document.styles[style_name]
    paragraph_properties = style._element.pPr
    run_properties = style._element.rPr
    spacing = paragraph_properties.find(qn("w:spacing"))
    color = run_properties.find(qn("w:color"))
    size = run_properties.find(qn("w:sz"))
    fonts = run_properties.find(qn("w:rFonts"))
    return {
        "before": spacing.get(qn("w:before")) if spacing is not None else None,
        "after": spacing.get(qn("w:after")) if spacing is not None else None,
        "line": spacing.get(qn("w:line")) if spacing is not None else None,
        "color": color.get(qn("w:val")) if color is not None else None,
        "size": size.get(qn("w:val")) if size is not None else None,
        "font": fonts.get(qn("w:ascii")) if fonts is not None else None,
    }


def test_results_docx_contains_required_sections_and_scientific_details(
    report_en: Path,
) -> None:
    """Dropping a required result field would produce an incomplete manuscript section."""
    document = Document(report_en)
    text = "\n".join(node.text or "" for node in document.element.iter(qn("w:t")))

    assert document.paragraphs[0].style.name == "Heading 1"
    assert document.paragraphs[0].text == "Results"
    assert "95% CI" in text
    assert "Table 1" in text
    assert "Figure 1" in text
    assert "n = 11" in text
    assert "missing = 1" in text
    assert "Variable-specific denominators ranged from n = 11–12" in text
    assert "missing counts ranged from 0–1" in text
    assert "p = 0.068" in text
    assert "p < 0.001" not in text
    assert "Hedges' g = -1.138" in text
    assert "Reproducibility" in text
    assert "Plan version: 3" in text
    assert len(document.tables) >= 1
    assert len(document.inline_shapes) == 1


def test_results_docx_uses_exact_manuscript_style_and_table_geometry(
    report_en: Path,
) -> None:
    """Renderer defaults or autofit would make journal geometry unstable across Word builds."""
    document = Document(report_en)
    section = document.sections[0]

    assert section.page_width.twips == 12240
    assert section.page_height.twips == 15840
    assert section.top_margin.twips == 1440
    assert section.right_margin.twips == 1440
    assert section.bottom_margin.twips == 1440
    assert section.left_margin.twips == 1440
    assert section.header_distance.twips == 708
    assert section.footer_distance.twips == 708
    assert not section.header.paragraphs[0].text
    assert not section.footer.paragraphs[0].text

    assert _style_properties(document, "Normal") == {
        "before": "0",
        "after": "120",
        "line": "264",
        "color": "000000",
        "size": "22",
        "font": "Calibri",
    }
    assert _style_properties(document, "Heading 1") == {
        "before": "320",
        "after": "160",
        "line": "264",
        "color": "183D3A",
        "size": "32",
        "font": "Calibri",
    }
    assert _style_properties(document, "Heading 2") == {
        "before": "240",
        "after": "120",
        "line": "264",
        "color": "183D3A",
        "size": "26",
        "font": "Calibri",
    }
    assert _style_properties(document, "Heading 3") == {
        "before": "160",
        "after": "80",
        "line": "264",
        "color": "2F635B",
        "size": "24",
        "font": "Calibri",
    }

    table = document.tables[0]
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    table_indent = properties.find(qn("w:tblInd"))
    table_layout = properties.find(qn("w:tblLayout"))
    assert table_width.get(qn("w:type")) == "dxa"
    assert table_width.get(qn("w:w")) == "9360"
    assert table_indent.get(qn("w:type")) == "dxa"
    assert table_indent.get(qn("w:w")) == "120"
    assert table_layout.get(qn("w:type")) == "fixed"
    widths = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.gridCol_lst]
    assert sum(widths) == 9360
    for row in table.rows:
        assert row.height is None
        assert [cell.width.twips for cell in row.cells] == widths
    assert table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None
    cell_margins = properties.find(qn("w:tblCellMar"))
    assert {
        child.tag.rsplit("}", 1)[-1]: child.get(qn("w:w")) for child in cell_margins
    } == {"top": "80", "bottom": "80", "start": "120", "end": "120"}


def test_results_docx_rebuilds_value_free_caption_alt_text_and_neutral_metadata(
    report_en: Path, report_inputs
) -> None:
    """Report-owned figure prose must remain meaningful without source labels or metadata."""
    _project, _brief, _plan, _bundle, figures = report_inputs
    document = Document(report_en)
    visible_text = "\n".join(node.text or "" for node in document.element.iter(qn("w:t")))
    descriptions = [
        node.get("descr")
        for node in document.element.iter(f"{{{WP_NS}}}docPr")
    ]

    expected_caption = (
        "Figure 1. Distribution for the planned outcome by planned exposure group; "
        "points show individual complete-case observations (n = 11)."
    )
    expected_alt = (
        "Group comparison figure showing individual complete-case observations and "
        "distributions for the planned outcome across planned exposure groups (n = 11)."
    )
    assert expected_caption in visible_text
    assert descriptions == [expected_alt]
    assert figures[0].caption not in visible_text
    assert figures[0].alt_text not in descriptions
    assert document.core_properties.author == "BioStat Studio"
    assert document.core_properties.last_modified_by == "BioStat Studio"

    payload = _complete_docx_payload(report_en)
    assert b"P001" not in payload
    assert str(FIXTURE_PATH).encode() not in payload
    assert b"/Users/" not in payload
    with ZipFile(report_en) as archive:
        document_xml = archive.read("word/document.xml")
    assert not re.search(rb"(?i)\b(?:todo|tbd|placeholder)\b|\{\{", document_xml)


@pytest.mark.parametrize("field", ["caption", "alt_text"])
@pytest.mark.parametrize(
    "unsafe_value",
    [
        "/private/tmp/patient-023.csv",
        "/Volumes/Clinic/patient-024.xlsx",
        "file:///private/tmp/patient-025",
        "https://example.test/patient-026",
        "line one\npatient-027",
        "patient-028 free-form note",
        "patient-029\x00control",
    ],
)
def test_figure_prose_injections_are_rebuilt_and_absent_from_complete_zip(
    tmp_path: Path,
    report_inputs,
    field: str,
    unsafe_value: str,
) -> None:
    """Caption and alt source fields must not become DOCX privacy channels."""
    project, brief, plan, bundle, figures = report_inputs
    replacement = (
        f"Figure 1. {unsafe_value}" if field == "caption" else unsafe_value
    )
    injected = replace(figures[0], **{field: replacement})
    destination = build_results_docx(
        project,
        brief,
        plan,
        bundle,
        [injected],
        "en",
        tmp_path / "safe-figure-prose.docx",
    )
    payload = _complete_docx_payload(destination)
    patient_token = re.search(r"patient-[0-9]+", unsafe_value)

    assert patient_token is not None
    assert patient_token.group().encode() not in payload
    assert unsafe_value.encode() not in payload
    assert "planned outcome" in _document_text(destination)
    descriptions = [
        node.get("descr")
        for node in Document(destination).element.iter(f"{{{WP_NS}}}docPr")
    ]
    assert descriptions == [
        "Group comparison figure showing individual complete-case observations and "
        "distributions for the planned outcome across planned exposure groups (n = 11)."
    ]


def test_turkish_table_uses_a_nonwrapping_not_applicable_token(
    tmp_path: Path, report_inputs
) -> None:
    """The Turkish not-applicable marker must not split one letter onto a new line."""
    project, brief, plan, bundle, _figures = report_inputs
    destination = build_results_docx(
        project, brief, plan, bundle, [], "tr", tmp_path / "table-tr.docx"
    )
    first_result = Document(destination).tables[0].rows[1]

    assert first_result.cells[1].text == "11–12"
    assert first_result.cells[2].text == "0–1"
    assert first_result.cells[5].text == "—"
    assert first_result.cells[6].text == "—"
    assert "Uygulanamaz" not in "\n".join(
        cell.text for row in Document(destination).tables[0].rows for cell in row.cells
    )


def test_warning_catalog_localizes_meaning_impact_and_action_without_raw_codes(
    tmp_path: Path, report_inputs
) -> None:
    """Warning codes are internal data and must become deterministic clinical copy."""
    project, brief, plan, bundle, _figures = report_inputs
    english = build_results_docx(
        project, brief, plan, bundle, [], "en", tmp_path / "warning-en.docx"
    )
    turkish = build_results_docx(
        project, brief, plan, bundle, [], "tr", tmp_path / "warning-tr.docx"
    )

    assert (
        "Meaning: analysis assumptions need review. Impact: unmet assumptions can make "
        "results unreliable. Action: review diagnostics before interpretation."
    ) in _document_text(english)
    assert (
        "Anlam: analiz varsayımları gözden geçirilmelidir. Etki: karşılanmayan "
        "varsayımlar sonuçları güvensiz kılabilir. Eylem: yorumlamadan önce tanıları "
        "inceleyin."
    ) in _document_text(turkish)
    assert b"review_assumptions" not in _complete_docx_payload(english)
    assert b"review_assumptions" not in _complete_docx_payload(turkish)


@pytest.mark.parametrize(
    "unsafe_warning",
    [
        "/private/tmp/patient-007.csv",
        "/Volumes/Clinic/patient-008.xlsx",
        "file:///private/tmp/patient-009",
        "https://example.test/patient-010",
        "line one\npatient-011",
        "patient-012 free-form note",
    ],
)
def test_unknown_warning_content_is_replaced_and_absent_from_complete_zip(
    tmp_path: Path, report_inputs, unsafe_warning: str
) -> None:
    """Unknown warning content must never become an exfiltration channel."""
    project, brief, plan, bundle, _figures = report_inputs
    plan.warnings = [unsafe_warning]
    destination = build_results_docx(
        project, brief, plan, bundle, [], "en", tmp_path / "redacted-warning.docx"
    )

    assert (
        "Meaning: an unrecognized analysis warning was recorded. Impact: the specific "
        "issue cannot be safely described in this report. Action: review the validated "
        "analysis log before interpretation."
    ) in _document_text(destination)
    assert unsafe_warning.encode() not in _complete_docx_payload(destination)


def test_structured_warning_ignores_untrusted_detail_fields(
    tmp_path: Path, report_inputs
) -> None:
    """Only an allowlisted structured warning code may cross the report boundary."""
    project, brief, plan, bundle, _figures = report_inputs
    plan.warnings = []
    bundle.warnings = [
        {
            "code": "library_warning",
            "method": "/private/tmp/patient-013",
            "category": "https://example.test/patient-014",
        }
    ]
    destination = build_results_docx(
        project, brief, plan, bundle, [], "en", tmp_path / "structured-warning.docx"
    )
    payload = _complete_docx_payload(destination)

    assert "Meaning: the statistical library emitted a warning." in _document_text(destination)
    assert b"/private/tmp/patient-013" not in payload
    assert b"https://example.test/patient-014" not in payload
    assert b"library_warning" not in payload


@pytest.mark.parametrize(
    ("field", "unsafe_value"),
    [
        ("exclusions", "/private/tmp/patient-015.csv"),
        ("exclusions", "/Volumes/Clinic/patient-016.xlsx"),
        ("exclusions", "file:///private/tmp/patient-017"),
        ("transformations", "https://example.test/patient-018"),
        ("transformations", "line one\npatient-019"),
        ("transformations", "patient-020 free-form note"),
        ("reproducibility", "3.9.6patient-021"),
        ("reproducibility", "3.9.6\x00patient-022"),
    ],
)
def test_unrecognized_provenance_is_rejected_before_docx_creation(
    tmp_path: Path, report_inputs, field: str, unsafe_value: str
) -> None:
    """Provenance accepts only recognized codes and safe version/count tokens."""
    project, brief, plan, bundle, _figures = report_inputs
    if field == "reproducibility":
        bundle.reproducibility["python"] = unsafe_value
    else:
        setattr(bundle.provenance, field, [unsafe_value])
    destination = tmp_path / "unsafe-provenance.docx"

    with pytest.raises(ValueError, match="unsafe_reproducibility_metadata"):
        build_results_docx(project, brief, plan, bundle, [], "en", destination)

    assert not destination.exists()


@pytest.mark.parametrize(
    ("value", "expected"),
    [(0.0009, "p < 0.001"), (0.001, "p = 0.001"), (0.068026, "p = 0.068")],
)
def test_p_value_formatting_uses_exact_values_unless_below_threshold(
    value: float, expected: str
) -> None:
    """Rounding a small exact p-value to zero would overstate the result."""
    assert format_p_value(value) == expected


def test_paired_report_uses_pairs_as_the_analysis_denominator(tmp_path: Path) -> None:
    """Reporting paired rows instead of complete pairs would double the scientific n."""
    frame = pd.DataFrame(
        {
            "pair_id": ["p1", "p1", "p2", "p2", "p3", "p3"],
            "condition": pd.Categorical(
                ["after", "before"] * 3,
                categories=["after", "before"],
                ordered=True,
            ),
            "score": [8.0, 5.0, 7.0, 6.0, float("nan"), 4.0],
        }
    )
    item = _item("primary_outcome", "paired_t_test", ["score", "condition", "pair_id"])
    plan = AnalysisPlan(items=[item])
    bundle = run_plan(frame, plan)
    brief = StudyBrief(
        title="Paired outcomes",
        question="Are paired outcome measurements associated with condition?",
        hypothesis="Paired outcome measurements differ by condition.",
        design="repeated",
        outcome_variables=["score"],
        exposure_variables=["condition"],
        pair_id_variable="pair_id",
    )

    destination = build_results_docx(
        LocalProject(tmp_path / "paired.biostat"),
        brief,
        plan,
        bundle,
        [],
        "en",
        tmp_path / "paired.docx",
    )
    text = "\n".join(
        node.text or "" for node in Document(destination).element.iter(qn("w:t"))
    )

    assert "n = 2 analysis units" in text
    assert "n = 4 analysis units" not in text

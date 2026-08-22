import json
import os
import subprocess
import sys
from pathlib import Path
from threading import Event

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from biostat_service.app import create_app


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setenv("BIOSTAT_SESSION_TOKEN", "test-token")
    with TestClient(create_app()) as test_client:
        yield test_client


def test_health_is_available_on_loopback(client):
    response = client.get("/health")

    assert response.json() == {"status": "ok", "service": "biostat-analysis", "api": 1}


def test_create_app_rejects_missing_session_token(monkeypatch):
    monkeypatch.delenv("BIOSTAT_SESSION_TOKEN", raising=False)

    with pytest.raises(RuntimeError, match="BIOSTAT_SESSION_TOKEN must be set"):
        create_app()


def test_create_app_rejects_empty_session_token(monkeypatch):
    monkeypatch.setenv("BIOSTAT_SESSION_TOKEN", "")

    with pytest.raises(RuntimeError, match="BIOSTAT_SESSION_TOKEN must be set"):
        create_app()


def test_v1_rejects_missing_token(client):
    response = client.get("/v1/session")

    assert response.status_code == 401


@pytest.mark.parametrize(
    ("method", "path"),
    [
        ("post", "/v1/data/profile"),
        ("post", "/v1/projects"),
        ("post", "/v1/projects/11111111-1111-4111-8111-111111111111/data-approval"),
        ("post", "/v1/plans"),
        ("post", "/v1/plans/approval"),
        ("post", "/v1/jobs"),
        ("get", "/v1/jobs/11111111-1111-4111-8111-111111111111"),
        ("post", "/v1/jobs/11111111-1111-4111-8111-111111111111/cancel"),
        ("post", "/v1/reports"),
    ],
)
def test_every_v1_capability_rejects_missing_token_before_body_validation(
    client, method: str, path: str
):
    response = client.post(path, json={}) if method == "post" else client.get(path)
    assert response.status_code == 401


def test_v1_rejects_raw_unprefixed_token(client):
    response = client.get("/v1/session", headers={"Authorization": "test-token"})

    assert response.status_code == 401


@pytest.mark.parametrize("authorization", ["Bearer", "Basic test-token", "BearerX test-token"])
def test_v1_rejects_malformed_authorization_scheme(client, authorization):
    response = client.get("/v1/session", headers={"Authorization": authorization})

    assert response.status_code == 401


def test_v1_rejects_wrong_bearer_token(client):
    response = client.get("/v1/session", headers={"Authorization": "Bearer wrong-token"})

    assert response.status_code == 401


def test_v1_accepts_session_token(client):
    response = client.get("/v1/session", headers={"Authorization": "Bearer test-token"})

    assert response.status_code == 200


def test_planning_requires_an_audited_data_structure_approval(client, tmp_path: Path):
    workbook_path = tmp_path / "approval.xlsx"
    workbook = Workbook()
    workbook.active.append(["group", "outcome"])
    workbook.active.append(["control", 1])
    workbook.active.append(["treated", 2])
    workbook.save(workbook_path)
    headers = {"Authorization": "Bearer test-token"}
    brief = {
        "title": "Approval boundary",
        "question": "Is the outcome different between groups?",
        "hypothesis": "The groups have different outcomes.",
        "design": "cohort",
        "outcome_variables": ["outcome"],
        "exposure_variables": ["group"],
    }
    created = client.post(
        "/v1/projects",
        headers=headers,
        json={"source_path": str(workbook_path), "project_root": str(tmp_path / "approval.biostat"), "brief": brief},
    )
    project_id = created.json()["id"]
    context = next(iter(client.app.state.projects.values()))
    assert context.approved_roles is None

    blocked = client.post("/v1/plans", headers=headers, json={"project_id": project_id})
    assert blocked.status_code == 409
    assert blocked.json() == {"detail": "data_structure_approval_required"}

    approved = client.post(
        f"/v1/projects/{project_id}/data-approval", headers=headers, json={}
    )
    assert approved.json() == {"approved": True}
    assert set(context.approved_roles) == {"group", "outcome"}
    assert all(role.confirmed for role in context.approved_roles.values())
    assert client.post("/v1/plans", headers=headers, json={"project_id": project_id}).status_code == 200
    audit = (tmp_path / "approval.biostat" / "audit.jsonl").read_text(encoding="utf-8")
    assert audit.count('"type": "data_structure_approved"') == 1


def test_new_plan_revision_invalidates_approval_and_rejects_stale_execution(client, tmp_path: Path):
    workbook_path = tmp_path / "revision.xlsx"
    workbook = Workbook()
    workbook.active.append(["group", "outcome"])
    workbook.active.append(["control", 1])
    workbook.active.append(["treated", 2])
    workbook.save(workbook_path)
    headers = {"Authorization": "Bearer test-token"}
    brief = {
        "title": "Revision binding",
        "question": "Is the outcome different between groups?",
        "hypothesis": "The groups have different outcomes.",
        "design": "cohort",
        "outcome_variables": ["outcome"],
        "exposure_variables": ["group"],
    }
    created = client.post(
        "/v1/projects", headers=headers,
        json={"source_path": str(workbook_path), "project_root": str(tmp_path / "revision.biostat"), "brief": brief},
    ).json()
    project_id = created["id"]
    client.post(f"/v1/projects/{project_id}/data-approval", headers=headers, json={})
    first = client.post("/v1/plans", headers=headers, json={"project_id": project_id}).json()
    assert first["revision"]
    assert len(first["digest"]) == 64
    assert client.post(
        "/v1/plans/approval", headers=headers,
        json={"project_id": project_id, "revision": first["revision"], "digest": first["digest"]},
    ).json() == {"approved": True, "revision": first["revision"], "digest": first["digest"]}

    second = client.post("/v1/plans", headers=headers, json={"project_id": project_id}).json()
    assert second["version"] == first["version"] == 1
    assert second["digest"] == first["digest"]
    assert second["revision"] != first["revision"]
    stale = client.post(
        "/v1/jobs", headers=headers,
        json={"project_id": project_id, "approved_plan_revision": first["revision"], "approved_plan_digest": first["digest"]},
    )
    assert stale.status_code == 409
    assert stale.json() == {"detail": "approved_plan_required"}


def test_validation_errors_never_echo_paths_uris_newlines_or_brief_text(client):
    headers = {"Authorization": "Bearer test-token"}
    hostile_path = "/Users/research/patient-001\nhttps://secret.example/study.xlsx"
    patient_text = "patient-001 has a rare diagnosis and must stay private"
    response = client.post(
        "/v1/projects",
        headers=headers,
        json={
            "source_path": hostile_path,
            "brief": {
                "title": "",
                "question": patient_text,
                "hypothesis": "x",
                "design": "cohort",
                "outcome_variables": [],
            },
        },
    )

    assert response.status_code == 422
    assert response.json() == {
        "error": {
            "code": "validation_error",
            "message": "Request validation failed.",
            "fields": [
                {"field": "brief.title", "code": "string_too_short"},
                {"field": "brief.hypothesis", "code": "string_too_short"},
                {"field": "brief.outcome_variables", "code": "too_short"},
            ],
        }
    }
    assert hostile_path not in response.text
    assert patient_text not in response.text
    assert "patient-001" not in response.text
    assert "secret.example" not in response.text


def test_report_uses_the_immutable_output_of_the_requested_job(client, tmp_path: Path, monkeypatch):
    workbook_path = tmp_path / "two-jobs.xlsx"
    workbook = Workbook()
    workbook.active.append(["group", "outcome"])
    for index in range(12):
        workbook.active.append(["control" if index < 6 else "treated", index + 1])
    workbook.save(workbook_path)
    headers = {"Authorization": "Bearer test-token"}
    brief = {
        "title": "Two jobs",
        "question": "Is the outcome different between groups?",
        "hypothesis": "The groups have different outcomes.",
        "design": "cohort",
        "outcome_variables": ["outcome"],
        "exposure_variables": ["group"],
    }
    project_id = client.post(
        "/v1/projects", headers=headers,
        json={"source_path": str(workbook_path), "project_root": str(tmp_path / "two-jobs.biostat"), "brief": brief},
    ).json()["id"]
    client.post(f"/v1/projects/{project_id}/data-approval", headers=headers, json={})

    def run_latest_plan():
        plan = client.post("/v1/plans", headers=headers, json={"project_id": project_id}).json()
        client.post(
            "/v1/plans/approval", headers=headers,
            json={"project_id": project_id, "revision": plan["revision"], "digest": plan["digest"]},
        )
        queued = client.post(
            "/v1/jobs", headers=headers,
            json={"project_id": project_id, "approved_plan_revision": plan["revision"], "approved_plan_digest": plan["digest"]},
        ).json()
        for _ in range(100):
            state = client.get(f"/v1/jobs/{queued['id']}", headers=headers).json()
            if state["status"] in {"completed", "failed", "cancelled"}:
                break
        assert state["status"] == "completed"
        return queued["id"]

    first_job = run_latest_plan()
    context = next(iter(client.app.state.projects.values()))
    first_bundle = context.bundle
    second_job = run_latest_plan()
    assert second_job != first_job
    assert context.bundle is not first_bundle

    exported_bundles = []
    def capture_report(_project, _brief, _plan, bundle, _figures, _language, destination):
        exported_bundles.append(bundle)
        Path(destination).write_bytes(b"docx fixture")

    monkeypatch.setattr("biostat_service.app.build_results_docx", capture_report)
    response = client.post(
        "/v1/reports", headers=headers,
        json={"project_id": project_id, "job_id": first_job, "language": "en", "destination": str(tmp_path / "first.docx")},
    )
    assert response.status_code == 200
    assert len(exported_bundles) == 1
    assert exported_bundles[0] is first_bundle


def test_cancelled_pipeline_cleans_staging_and_never_becomes_reportable(client, tmp_path: Path, monkeypatch):
    workbook_path = tmp_path / "cancel.xlsx"
    workbook = Workbook()
    workbook.active.append(["group", "outcome"])
    for index in range(12):
        workbook.active.append(["control" if index < 6 else "treated", index + 1])
    workbook.save(workbook_path)
    headers = {"Authorization": "Bearer test-token"}
    brief = {
        "title": "Cancellation",
        "question": "Is the outcome different between groups?",
        "hypothesis": "The groups have different outcomes.",
        "design": "cohort",
        "outcome_variables": ["outcome"],
        "exposure_variables": ["group"],
    }
    root = tmp_path / "cancel.biostat"
    project_id = client.post(
        "/v1/projects", headers=headers,
        json={"source_path": str(workbook_path), "project_root": str(root), "brief": brief},
    ).json()["id"]
    client.post(f"/v1/projects/{project_id}/data-approval", headers=headers, json={})
    plan = client.post("/v1/plans", headers=headers, json={"project_id": project_id}).json()
    client.post(
        "/v1/plans/approval", headers=headers,
        json={"project_id": project_id, "revision": plan["revision"], "digest": plan["digest"]},
    )

    figures_started = Event()
    release_figures = Event()
    def blocking_figures(_frame, _plan, _bundle, output_dir, _language):
        Path(output_dir).mkdir(parents=True)
        (Path(output_dir) / "partial.png").write_bytes(b"partial")
        figures_started.set()
        release_figures.wait(timeout=2)
        return []

    monkeypatch.setattr("biostat_service.app.build_figures", blocking_figures)
    queued = client.post(
        "/v1/jobs", headers=headers,
        json={"project_id": project_id, "approved_plan_revision": plan["revision"], "approved_plan_digest": plan["digest"]},
    ).json()
    assert figures_started.wait(timeout=2)
    cancelled = client.post(f"/v1/jobs/{queued['id']}/cancel", headers=headers, json={})
    assert cancelled.json()["status"] == "cancelling"
    release_figures.set()
    for _ in range(100):
        state = client.get(f"/v1/jobs/{queued['id']}", headers=headers).json()
        if state["status"] == "cancelled":
            break

    assert state["status"] == "cancelled"
    assert state["result"] is None
    assert '"type": "analysis_completed"' not in (root / "audit.jsonl").read_text(encoding="utf-8")
    assert not list((root / "artifacts").glob(".staging-*"))
    assert next(iter(client.app.state.projects.values())).job_outputs == {}
    report = client.post(
        "/v1/reports", headers=headers,
        json={"project_id": project_id, "job_id": queued["id"], "language": "en", "destination": str(tmp_path / "cancelled.docx")},
    )
    assert report.status_code == 409
    assert not (tmp_path / "cancelled.docx").exists()


def test_authenticated_project_plan_job_and_report_routes_keep_values_out_of_responses(
    client, tmp_path: Path
):
    """The entire local workflow publishes metadata/results, never raw workbook values."""
    workbook_path = tmp_path / "study.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["group", "outcome"])
    for index in range(12):
        sheet.append(["control" if index < 6 else "treated", 100 + index])
    workbook.save(workbook_path)
    brief = {
        "title": "Treatment comparison",
        "question": "Does treatment change the measured outcome?",
        "hypothesis": "Treatment is associated with a different outcome.",
        "design": "cohort",
        "outcome_variables": ["outcome"],
        "exposure_variables": ["group"],
        "language": "en",
    }
    headers = {"Authorization": "Bearer test-token"}

    created = client.post(
        "/v1/projects",
        headers=headers,
        json={"source_path": str(workbook_path), "project_root": str(tmp_path / "study.biostat"), "brief": brief},
    )
    assert created.status_code == 200
    payload = created.json()
    assert payload["profile"]["rows"] == 12
    assert "control" not in created.text
    assert str(workbook_path) not in created.text

    approved_data = client.post(
        f"/v1/projects/{payload['id']}/data-approval", headers=headers, json={}
    )
    assert approved_data.status_code == 200

    planned = client.post("/v1/plans", headers=headers, json={"project_id": payload["id"]})
    assert planned.status_code == 200
    assert planned.json()["blocking_errors"] == []
    approved_plan = client.post(
        "/v1/plans/approval",
        headers=headers,
        json={
            "project_id": payload["id"],
            "revision": planned.json()["revision"],
            "digest": planned.json()["digest"],
        },
    )
    assert approved_plan.status_code == 200

    queued = client.post(
        "/v1/jobs",
        headers=headers,
        json={
            "project_id": payload["id"],
            "approved_plan_revision": planned.json()["revision"],
            "approved_plan_digest": planned.json()["digest"],
        },
    )
    assert queued.status_code == 200
    job_id = queued.json()["id"]
    for _ in range(100):
        status = client.get(f"/v1/jobs/{job_id}", headers=headers)
        assert status.status_code == 200
        if status.json()["status"] in {"completed", "failed", "cancelled"}:
            break
    assert status.json()["status"] == "completed"
    assert status.json()["result"]["results"]
    assert "control" not in status.text

    audit_text = (tmp_path / "study.biostat" / "audit.jsonl").read_text(encoding="utf-8")
    assert '"type": "analysis_completed"' in audit_text
    assert '"data_fingerprint"' in audit_text
    assert "control" not in audit_text

    destination = tmp_path / "results.docx"
    exported = client.post(
        "/v1/reports",
        headers=headers,
        json={"project_id": payload["id"], "job_id": job_id, "language": "en", "destination": str(destination)},
    )
    assert exported.status_code == 200
    assert exported.json() == {"saved": True, "filename": "results.docx"}
    assert destination.exists()


def test_v1_rejects_non_loopback_client(monkeypatch):
    monkeypatch.setenv("BIOSTAT_SESSION_TOKEN", "test-token")
    with TestClient(create_app(), client=("192.0.2.1", 5000)) as remote_client:
        response = remote_client.get(
            "/v1/session", headers={"Authorization": "Bearer test-token"}
        )

    assert response.status_code == 403
    assert response.json() == {"detail": "loopback_only"}


def test_launcher_emits_only_port_and_api_readiness():
    process = subprocess.Popen(
        [sys.executable, "-m", "biostat_service.app", "--port", "0"],
        env={**os.environ, "BIOSTAT_SESSION_TOKEN": "test-token"},
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    try:
        assert process.stdout is not None
        readiness = json.loads(process.stdout.readline())
        assert readiness.keys() == {"port", "api"}
        assert isinstance(readiness["port"], int)
        assert readiness["port"] > 0
        assert readiness["api"] == 1
    finally:
        process.terminate()
        process.wait(timeout=5)

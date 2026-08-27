"""Privacy contract for the real-file methodology acceptance runner."""

from __future__ import annotations

import importlib.util
import json
import sys
from pathlib import Path
from types import ModuleType

from openpyxl import Workbook

from biostat_service.extractors.contracts import BriefProposal, Proposal, RoleProposal
from biostat_service.extractors.engine import (
    EngineStatus,
    ExtractionRun,
    VariableRun,
)


def _load_runner() -> ModuleType:
    path = Path(__file__).parents[1] / "scripts" / "accept-methodology.py"
    spec = importlib.util.spec_from_file_location("accept_methodology", path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class StubEngine:
    def extract(self, document):
        evidence = "The primary outcome was clinical deterioration."
        return ExtractionRun(
            BriefProposal(
                title=Proposal("Clinical deterioration", 0.79, "local:test", evidence, 8),
                question=Proposal("Does notification predict deterioration?", 0.79, "local:test", evidence, 8),
                hypothesis=Proposal("Notification predicts deterioration.", 0.79, "local:test", evidence, 8),
                design=Proposal("cohort", 0.79, "local:test", "cohort", 0),
                outcome_concepts=(
                    Proposal("clinical deterioration", 0.79, "local:test", evidence, 8),
                ),
                exposure_concepts=(
                    Proposal("notification", 0.79, "local:test", evidence, 8),
                ),
            ),
            EngineStatus("local:test", "local:test"),
        )

    def variables(self, document, columns):
        evidence = "The primary outcome was clinical deterioration."
        return VariableRun(
            (
                RoleProposal(
                    "outcome_column",
                    role=Proposal("outcome", 0.79, "local:test", evidence, 8),
                ),
                RoleProposal(
                    "exposure_column",
                    role=Proposal("exposure", 0.79, "local:test", evidence, 8),
                ),
            ),
            EngineStatus("local:test", "local:test"),
        )


def test_acceptance_summary_never_serializes_paths_evidence_or_cell_values(
    tmp_path: Path,
) -> None:
    runner = _load_runner()
    document = tmp_path / "methods.docx"
    from docx import Document

    source = Document()
    source.add_paragraph("cohort")
    source.add_paragraph("The primary outcome was clinical deterioration.")
    source.save(document)

    workbook_path = tmp_path / "research.xlsx"
    workbook = Workbook()
    workbook.active.append(["outcome_column", "exposure_column"])
    workbook.active.append(["PRIVATE_CELL_SENTINEL", 1])
    workbook.save(workbook_path)

    summary = runner.acceptance_summary(document, workbook_path, engine=StubEngine())
    rendered = json.dumps(summary, ensure_ascii=False)

    assert summary["document"]["name"] == "methods.docx"
    assert summary["workbook"] == {
        "name": "research.xlsx",
        "sheet": "Sheet",
        "rows": 1,
        "columns": 2,
    }
    assert summary["engine"]["brief"]["used"] == "local:test"
    assert summary["brief"]["outcomes"] == ["clinical deterioration"]
    assert summary["roles"] == [
        {
            "column": "outcome_column",
            "role": "outcome",
            "kind_correction": None,
            "confidence": 0.79,
            "source": "local:test",
        },
        {
            "column": "exposure_column",
            "role": "exposure",
            "kind_correction": None,
            "confidence": 0.79,
            "source": "local:test",
        },
    ]
    assert str(tmp_path) not in rendered
    assert "The primary outcome" not in rendered
    assert "PRIVATE_CELL_SENTINEL" not in rendered
    assert "evidence" not in rendered

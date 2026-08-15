"""Deterministic bilingual Word manuscript sections for completed analyses."""

from __future__ import annotations

from datetime import datetime, timezone
import math
from pathlib import Path
import re
from typing import Literal, Sequence

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips

from biostat_service.analyses import AnalysisBundle
from biostat_service.contracts import AnalysisPlan, AnalysisResult, PlanItem, StudyBrief
from biostat_service.projects import LocalProject
from biostat_service.visuals import FigureArtifact


Language = Literal["en", "tr"]
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
TABLE_CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}
TABLE_COLUMN_WIDTHS_DXA = (2500, 700, 950, 1100, 1500, 1200, 1410)
FINGERPRINT_PREFIX_LENGTH = 12
HEADING_COLOR = "183D3A"
HEADING_3_COLOR = "2F635B"
TABLE_HEADER_FILL = "E8EFE9"
TERRACOTTA = "D66E47"
TABLE_BORDER_COLOR = "AABAB8"
MAX_FIGURE_WIDTH_INCHES = 5.0
SAFE_FINGERPRINT = re.compile(r"[0-9a-f]{64}\Z")
SAFE_VERSION = re.compile(
    r"(?:0|[1-9][0-9]*)(?:\.(?:0|[1-9][0-9]*)){1,3}\Z"
)
SAFE_EXCLUSION = re.compile(
    r"(?:complete_case|paired_complete_case):"
    r"input=(?P<input>[0-9]+);used=(?P<used>[0-9]+);missing=(?P<missing>[0-9]+)\Z"
)
SAFE_SOFTWARE_COMPONENTS = frozenset(
    {"python", "numpy", "pandas", "scipy", "statsmodels"}
)
SAFE_TRANSFORMATIONS = frozenset({"deterministic_complete_case_execution"})
OBSERVATIONAL_DESIGNS = frozenset({"cross_sectional", "cohort", "case_control"})
FIXED_METADATA_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)


RESULT_LABELS = {
    "en": {
        "heading": "Results",
        "findings": "Statistical findings",
        "table": "Table",
        "table_caption": "Statistical results.",
        "figure": "Figure",
        "figures": "Figures",
        "warnings": "Material warnings",
        "reproducibility": "Reproducibility",
        "analysis": "Analysis",
        "missing": "Missing",
        "estimate": "Estimate",
        "ci": "95% CI",
        "p_value": "p value",
        "effect": "Effect size",
        "not_estimable": "Not estimable",
        "not_applicable": "Not applicable",
        "narrative": (
            "The {analysis} included n = {n} analysis units (missing = {missing}). "
            "The estimated {parameter} was {estimate} (95% CI [{lower}, {upper}]), "
            "{p_value}, with {effect_name} = {effect_value}."
        ),
        "descriptive_narrative": (
            "The descriptive analysis included n = {n} analysis units (missing = {missing}). "
            "The first numeric estimate was {estimate} (95% CI [{lower}, {upper}]); "
            "no hypothesis-test p value or standardized effect size was applicable."
        ),
        "observational_note": (
            "This observational analysis describes associations and does not establish causality."
        ),
        "table_note": (
            "Note. CI = confidence interval; n = analyzed units. "
            "Exact p values are shown when available."
        ),
        "plan_version": "Plan version: {value}",
        "fingerprint": "Source fingerprint prefix: {value}",
        "software": "Software versions: {value}",
        "exclusions": "Exclusions: {value}",
        "transformations": "Transformations: {value}",
        "random_seed": "Random seed: not used",
    },
    "tr": {
        "heading": "Bulgular",
        "findings": "İstatistiksel bulgular",
        "table": "Tablo",
        "table_caption": "İstatistiksel sonuçlar.",
        "figure": "Şekil",
        "figures": "Şekiller",
        "warnings": "Önemli uyarılar",
        "reproducibility": "Yeniden üretilebilirlik",
        "analysis": "Analiz",
        "missing": "Eksik",
        "estimate": "Tahmin",
        "ci": "%95 GA",
        "p_value": "p değeri",
        "effect": "Etki büyüklüğü",
        "not_estimable": "Tahmin edilemedi",
        "not_applicable": "Uygulanamaz",
        "narrative": (
            "{analysis} analizi n = {n} analiz birimi içerdi (eksik = {missing}). "
            "Tahmin edilen {parameter} {estimate} idi (%95 GA [{lower}, {upper}]), "
            "{p_value}; {effect_name} = {effect_value}."
        ),
        "descriptive_narrative": (
            "Betimsel analiz n = {n} analiz birimi içerdi (eksik = {missing}). "
            "İlk sayısal tahmin {estimate} idi (%95 GA [{lower}, {upper}]); "
            "hipotez testi p değeri ve standartlaştırılmış etki büyüklüğü uygulanamazdı."
        ),
        "observational_note": (
            "Bu gözlemsel analiz ilişkileri betimler ve nedensellik göstermez."
        ),
        "table_note": (
            "Not. GA = güven aralığı; n = analiz edilen birimler. "
            "Kesin p değerleri mevcut olduğunda gösterilir."
        ),
        "plan_version": "Plan sürümü: {value}",
        "fingerprint": "Kaynak parmak izi öneki: {value}",
        "software": "Yazılım sürümleri: {value}",
        "exclusions": "Dışlamalar: {value}",
        "transformations": "Dönüşümler: {value}",
        "random_seed": "Rastgelelik tohumu: kullanılmadı",
    },
}

WARNING_CATALOG = {
    "en": {
        "review_assumptions": (
            "Meaning: analysis assumptions need review. Impact: unmet assumptions can "
            "make results unreliable. Action: review diagnostics before interpretation."
        ),
        "data_profile:empty_column": (
            "Meaning: a planned variable has no observed values. Impact: analyses using "
            "that variable may not be estimable. Action: review data completeness and the "
            "confirmed analysis plan."
        ),
        "data_profile:mixed_types": (
            "Meaning: a planned variable contains mixed data types. Impact: coercion or "
            "coding choices may affect the analysis. Action: verify the variable type and "
            "coding before interpretation."
        ),
        "data_profile:non_finite_values": (
            "Meaning: a planned variable contains non-finite numeric values. Impact: "
            "affected rows may be excluded or the analysis may be invalid. Action: review "
            "the source data and exclusion counts."
        ),
        "data_profile:duplicated_identifier": (
            "Meaning: an identifier-labelled variable contains duplicates. Impact: unit "
            "or pair definitions may be ambiguous. Action: verify identifier uniqueness "
            "before interpretation."
        ),
        "data_profile:suspicious_identifier_leakage": (
            "Meaning: an identifier-labelled variable may contain identifying information. "
            "Impact: privacy could be compromised if it is included in analysis artifacts. "
            "Action: confirm de-identification and exclude identifiers from reporting."
        ),
        "zero_cell_odds_ratio_corrected": (
            "Meaning: a zero contingency-table cell required a continuity correction. "
            "Impact: the odds-ratio estimate is correction-dependent. Action: interpret "
            "the estimate with the exact test and cell sparsity in mind."
        ),
        "fisher_odds_ratio_non_finite": (
            "Meaning: the Fisher-test odds ratio was not finite. Impact: the effect estimate "
            "cannot be interpreted as a finite ratio. Action: review the contingency table "
            "and report the exact test result with this limitation."
        ),
        "library_warning": (
            "Meaning: the statistical library emitted a warning. Impact: numerical "
            "stability or model assumptions may affect the result. Action: review the "
            "validated diagnostics before interpretation."
        ),
        "unknown": (
            "Meaning: an unrecognized analysis warning was recorded. Impact: the specific "
            "issue cannot be safely described in this report. Action: review the validated "
            "analysis log before interpretation."
        ),
    },
    "tr": {
        "review_assumptions": (
            "Anlam: analiz varsayımları gözden geçirilmelidir. Etki: karşılanmayan "
            "varsayımlar sonuçları güvensiz kılabilir. Eylem: yorumlamadan önce "
            "tanıları inceleyin."
        ),
        "data_profile:empty_column": (
            "Anlam: planlanan bir değişkende gözlenen değer yoktur. Etki: bu değişkeni "
            "kullanan analizler tahmin edilemeyebilir. Eylem: veri tamlığını ve onaylanmış "
            "analiz planını gözden geçirin."
        ),
        "data_profile:mixed_types": (
            "Anlam: planlanan bir değişken karma veri türleri içerir. Etki: dönüştürme "
            "veya kodlama seçimleri analizi etkileyebilir. Eylem: yorumlamadan önce değişken "
            "türünü ve kodlamayı doğrulayın."
        ),
        "data_profile:non_finite_values": (
            "Anlam: planlanan bir değişken sonlu olmayan sayısal değerler içerir. Etki: "
            "etkilenen satırlar dışlanabilir veya analiz geçersiz olabilir. Eylem: kaynak "
            "veriyi ve dışlama sayılarını gözden geçirin."
        ),
        "data_profile:duplicated_identifier": (
            "Anlam: tanımlayıcı olarak etiketlenen bir değişken yinelenen değerler içerir. "
            "Etki: birim veya çift tanımları belirsiz olabilir. Eylem: yorumlamadan önce "
            "tanımlayıcı benzersizliğini doğrulayın."
        ),
        "data_profile:suspicious_identifier_leakage": (
            "Anlam: tanımlayıcı olarak etiketlenen bir değişken kimlik bilgisi içerebilir. "
            "Etki: analiz yapıtlarına eklenirse gizlilik tehlikeye girebilir. Eylem: "
            "kimliksizleştirmeyi "
            "doğrulayın ve tanımlayıcıları raporlamadan dışlayın."
        ),
        "zero_cell_odds_ratio_corrected": (
            "Anlam: sıfır kontenjans tablosu hücresi süreklilik düzeltmesi gerektirdi. Etki: "
            "olasılık oranı tahmini düzeltmeye bağlıdır. Eylem: tahmini kesin test "
            "ve seyrek "
            "hücreler ile birlikte yorumlayın."
        ),
        "fisher_odds_ratio_non_finite": (
            "Anlam: Fisher testi olasılık oranı sonlu değildir. Etki: etki tahmini sonlu bir "
            "oran olarak yorumlanamaz. Eylem: kontenjans tablosunu gözden geçirin ve kesin test "
            "sonucunu bu sınırlılıkla raporlayın."
        ),
        "library_warning": (
            "Anlam: istatistik kütüphanesi bir uyarı verdi. Etki: sayısal kararlılık veya "
            "model varsayımları sonucu etkileyebilir. Eylem: yorumlamadan önce doğrulanmış "
            "tanıları gözden geçirin."
        ),
        "unknown": (
            "Anlam: tanınmayan bir analiz uyarısı kaydedildi. Etki: belirli sorun bu raporda "
            "güvenli biçimde açıklanamaz. Eylem: yorumlamadan önce doğrulanmış analiz "
            "günlüğünü inceleyin."
        ),
    },
}

METHOD_LABELS = {
    "en": {
        "descriptive_summary": ("descriptive analysis", "descriptive estimate"),
        "welch_t_test": ("Welch two-group comparison", "mean difference"),
        "paired_t_test": ("paired comparison", "paired mean difference"),
        "welch_anova": ("Welch omnibus comparison", "omnibus effect"),
        "chi_square_or_fisher": ("categorical association analysis", "odds ratio"),
        "pearson_or_spearman": ("correlation analysis", "correlation"),
        "linear_regression": ("linear regression", "regression coefficient"),
        "logistic_regression": ("logistic regression", "odds ratio"),
    },
    "tr": {
        "descriptive_summary": ("betimsel analiz", "betimsel tahmin"),
        "welch_t_test": ("Welch iki grup karşılaştırması", "ortalama farkı"),
        "paired_t_test": ("eşleştirilmiş karşılaştırma", "eşleştirilmiş ortalama farkı"),
        "welch_anova": ("Welch genel karşılaştırması", "genel etki"),
        "chi_square_or_fisher": ("kategorik ilişki analizi", "olasılık oranı"),
        "pearson_or_spearman": ("korelasyon analizi", "korelasyon"),
        "linear_regression": ("doğrusal regresyon", "regresyon katsayısı"),
        "logistic_regression": ("lojistik regresyon", "olasılık oranı"),
    },
}

EFFECT_LABELS = {
    "en": {
        "hedges_g": "Hedges' g",
        "paired_hedges_g_z": "paired Hedges' g-z",
        "welch_cohen_f_squared": "Welch-compatible f-squared",
        "odds_ratio": "odds ratio",
        "pearson_r": "Pearson's r",
        "regression_coefficient": "regression coefficient",
        "standardized_descriptive_estimate": "standardized descriptive estimate",
    },
    "tr": {
        "hedges_g": "Hedges' g",
        "paired_hedges_g_z": "eşleştirilmiş Hedges' g-z",
        "welch_cohen_f_squared": "Welch uyumlu f-kare",
        "odds_ratio": "olasılık oranı",
        "pearson_r": "Pearson r",
        "regression_coefficient": "regresyon katsayısı",
        "standardized_descriptive_estimate": "standartlaştırılmış betimsel tahmin",
    },
}


def format_p_value(value: float) -> str:
    """Format a valid p value without ever rounding a small value to zero."""
    numeric = float(value)
    if not math.isfinite(numeric) or not 0 <= numeric <= 1:
        raise ValueError("invalid_p_value")
    return "p < 0.001" if numeric < 0.001 else f"p = {numeric:.3f}"


def _format_number(value: float | int | None) -> str | None:
    if value is None:
        return None
    numeric = float(value)
    if not math.isfinite(numeric):
        raise ValueError("non_finite_report_value")
    if abs(numeric) < 0.0005:
        numeric = 0.0
    return f"{numeric:.3f}"


def _set_style_font(style, name: str, size: float, color: str, *, bold: bool | None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    run_properties = style._element.get_or_add_rPr()
    fonts = run_properties.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), name)


def _set_style_spacing(
    style,
    *,
    before: float,
    after: float,
    line_spacing: float,
    keep_with_next: bool = False,
) -> None:
    formatting = style.paragraph_format
    formatting.space_before = Pt(before)
    formatting.space_after = Pt(after)
    formatting.line_spacing = line_spacing
    formatting.keep_with_next = keep_with_next
    formatting.keep_together = keep_with_next


def _configure_styles(document: DocumentType) -> None:
    normal = document.styles["Normal"]
    _set_style_font(normal, "Calibri", 11, "000000", bold=False)
    _set_style_spacing(normal, before=0, after=6, line_spacing=1.10)

    heading_tokens = {
        "Heading 1": (16, HEADING_COLOR, 16, 8),
        "Heading 2": (13, HEADING_COLOR, 12, 6),
        "Heading 3": (12, HEADING_3_COLOR, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[name]
        _set_style_font(style, "Calibri", size, color, bold=True)
        _set_style_spacing(
            style,
            before=before,
            after=after,
            line_spacing=1.10,
            keep_with_next=True,
        )

    for name, size in (("Title", 20), ("Subtitle", 12)):
        style = document.styles[name]
        _set_style_font(style, "Calibri", size, HEADING_COLOR, bold=name == "Title")
        _set_style_spacing(style, before=0, after=6, line_spacing=1.10)

    caption = document.styles["Caption"]
    _set_style_font(caption, "Calibri", 9, TERRACOTTA, bold=False)
    _set_style_spacing(
        caption,
        before=4,
        after=4,
        line_spacing=1.0,
        keep_with_next=False,
    )


def _configure_page(document: DocumentType) -> None:
    for section in document.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Twips(708)
        section.footer_distance = Twips(708)
        for paragraph in (*section.header.paragraphs, *section.footer.paragraphs):
            paragraph.text = ""


def _configure_metadata(document: DocumentType) -> None:
    properties = document.core_properties
    properties.title = "BioStat Studio"
    properties.subject = "Statistical results"
    properties.author = "BioStat Studio"
    properties.last_modified_by = "BioStat Studio"
    properties.keywords = ""
    properties.comments = ""
    properties.category = ""
    properties.created = FIXED_METADATA_TIME
    properties.modified = FIXED_METADATA_TIME


def _counts(result: AnalysisResult) -> tuple[int, int]:
    raw = result.diagnostics.get("counts", {})
    used = raw.get("used", result.n)
    missing = raw.get("missing", 0)
    if type(used) is not int or type(missing) is not int or used < 0 or missing < 0:
        raise ValueError("invalid_result_counts")
    return result.n, missing


def _result_labels(item: PlanItem, language: Language) -> tuple[str, str]:
    try:
        return METHOD_LABELS[language][item.method]
    except KeyError as exc:
        raise ValueError(f"unsupported_report_method:{item.method}") from exc


def _effect_label(result: AnalysisResult, language: Language) -> str:
    return EFFECT_LABELS[language].get(result.effect_size.name, result.effect_size.name)


def _result_values(result: AnalysisResult, language: Language) -> dict[str, str]:
    labels = RESULT_LABELS[language]
    used, missing = _counts(result)
    lower = _format_number(result.confidence_interval.lower)
    upper = _format_number(result.confidence_interval.upper)
    estimate = _format_number(result.estimate)
    effect = _format_number(result.effect_size.value)
    return {
        "n": str(used),
        "missing": str(missing),
        "estimate": estimate or labels["not_estimable"],
        "lower": lower or labels["not_estimable"],
        "upper": upper or labels["not_estimable"],
        "ci": (
            f"[{lower}, {upper}]"
            if lower is not None and upper is not None
            else labels["not_estimable"]
        ),
        "p_value": (
            format_p_value(result.p_value)
            if result.p_value is not None
            else labels["not_applicable"]
        ),
        "effect_name": _effect_label(result, language),
        "effect_value": effect or labels["not_estimable"],
        "effect": (
            f"{_effect_label(result, language)} = {effect}"
            if effect is not None
            else labels["not_applicable"]
        ),
    }


def _add_results_narrative(
    document: DocumentType,
    brief: StudyBrief,
    plan: AnalysisPlan,
    bundle: AnalysisBundle,
    language: Language,
) -> None:
    labels = RESULT_LABELS[language]
    document.add_heading(labels["findings"], level=2)
    for item in plan.items:
        result = bundle.results[item.id]
        analysis, parameter = _result_labels(item, language)
        document.add_heading(analysis.capitalize(), level=3)
        values = _result_values(result, language)
        if item.method == "descriptive_summary":
            narrative = labels["descriptive_narrative"].format(**values)
        else:
            narrative = labels["narrative"].format(
                analysis=analysis,
                parameter=parameter,
                **values,
            )
        document.add_paragraph(narrative)

    if brief.design in OBSERVATIONAL_DESIGNS:
        paragraph = document.add_paragraph(labels["observational_note"])
        paragraph.paragraph_format.keep_together = True


def _set_table_cell_margins(table) -> None:
    properties = table._tbl.tblPr
    existing = properties.find(qn("w:tblCellMar"))
    if existing is not None:
        properties.remove(existing)
    margins = OxmlElement("w:tblCellMar")
    for side, width in TABLE_CELL_MARGINS_DXA.items():
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(width))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def _set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    existing = properties.find(qn("w:tblBorders"))
    if existing is not None:
        properties.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), TABLE_BORDER_COLOR)
        borders.append(border)
    properties.append(borders)


def _set_table_geometry(table, widths: Sequence[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA or len(widths) != len(table.columns):
        raise ValueError("invalid_table_geometry")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.insert(0, table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))

    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.insert(1, indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))

    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid_columns = list(table._tbl.tblGrid.gridCol_lst)
    for column, width in zip(grid_columns, widths):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Twips(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.get_or_add_tcW()
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(width))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(
                        qn("w:ascii"), "Calibri"
                    )

    _set_table_cell_margins(table)
    _set_table_borders(table)


def _mark_header_row(row) -> None:
    properties = row._tr.get_or_add_trPr()
    existing = properties.find(qn("w:tblHeader"))
    if existing is None:
        existing = OxmlElement("w:tblHeader")
        properties.append(existing)
    existing.set(qn("w:val"), "true")
    for cell in row.cells:
        shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            cell._tc.get_or_add_tcPr().append(shading)
        shading.set(qn("w:fill"), TABLE_HEADER_FILL)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(HEADING_COLOR)


def _add_result_table(
    document: DocumentType,
    plan: AnalysisPlan,
    bundle: AnalysisBundle,
    language: Language,
) -> None:
    labels = RESULT_LABELS[language]
    caption = document.add_paragraph(
        f"{labels['table']} 1. {labels['table_caption']}", style="Caption"
    )
    caption.paragraph_format.keep_with_next = True
    headers = (
        labels["analysis"],
        "n",
        labels["missing"],
        labels["estimate"],
        labels["ci"],
        labels["p_value"],
        labels["effect"],
    )
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text

    for item in plan.items:
        result = bundle.results[item.id]
        analysis, _parameter = _result_labels(item, language)
        values = _result_values(result, language)
        row_values = (
            analysis.capitalize(),
            values["n"],
            values["missing"],
            values["estimate"],
            values["ci"],
            values["p_value"],
            values["effect"],
        )
        for index, (cell, text) in enumerate(zip(table.add_row().cells, row_values)):
            cell.text = text
            cell.paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.LEFT if index in (0, 6) else WD_ALIGN_PARAGRAPH.CENTER
            )

    _set_table_geometry(table, TABLE_COLUMN_WIDTHS_DXA)
    _mark_header_row(table.rows[0])
    note = document.add_paragraph(labels["table_note"], style="Caption")
    note.paragraph_format.keep_together = True


def _add_figures(
    document: DocumentType,
    figures: Sequence[FigureArtifact],
    language: Language,
) -> None:
    if not figures:
        return
    labels = RESULT_LABELS[language]
    document.add_heading(labels["figures"], level=2)
    expected_prefix = f"{labels['figure']} "
    for number, artifact in enumerate(figures, start=1):
        png_path = Path(artifact.png_path)
        if artifact.dpi != 300 or png_path.suffix.lower() != ".png" or not png_path.is_file():
            raise ValueError("invalid_report_figure_png")
        if not artifact.caption.startswith(f"{expected_prefix}{number}."):
            raise ValueError("figure_caption_language_or_number_mismatch")
        if not artifact.alt_text.strip():
            raise ValueError("missing_figure_alt_text")
        picture = document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.keep_with_next = True
        picture.paragraph_format.keep_together = True
        run = picture.add_run()
        inline_shape = run.add_picture(
            str(png_path), width=Inches(min(artifact.width_inches, MAX_FIGURE_WIDTH_INCHES))
        )
        inline_shape._inline.docPr.set("descr", artifact.alt_text)
        inline_shape._inline.docPr.set("title", f"{labels['figure']} {number}")
        caption = document.add_paragraph(artifact.caption, style="Caption")
        caption.paragraph_format.keep_together = True


def _warning_catalog_key(warning: object) -> str:
    """Map a structured warning to an allowlisted key without retaining raw details."""
    if isinstance(warning, dict):
        return "library_warning" if warning.get("code") == "library_warning" else "unknown"
    if not isinstance(warning, str):
        return "unknown"
    if warning in WARNING_CATALOG["en"] and warning != "unknown":
        return warning
    if warning.startswith("data_profile:"):
        parts = warning.split(":", 2)
        key = f"data_profile:{parts[1]}" if len(parts) == 3 else "unknown"
        return key if key in WARNING_CATALOG["en"] else "unknown"
    if warning.startswith("library_warning:"):
        return "library_warning"
    return "unknown"


def _add_warnings(
    document: DocumentType,
    plan: AnalysisPlan,
    bundle: AnalysisBundle,
    language: Language,
) -> None:
    warnings: list[object] = [*plan.warnings, *bundle.warnings]
    for result in bundle.results.values():
        warnings.extend(result.warnings)
    if not warnings:
        return
    labels = RESULT_LABELS[language]
    document.add_heading(labels["warnings"], level=2)
    keys = dict.fromkeys(_warning_catalog_key(value) for value in warnings)
    for key in keys:
        paragraph = document.add_paragraph(WARNING_CATALOG[language][key])
        paragraph.paragraph_format.keep_together = True


def _safe_exclusions(values: Sequence[str], fallback: str) -> str:
    cleaned: list[str] = []
    for value in values:
        if not isinstance(value, str):
            raise ValueError("unsafe_reproducibility_metadata")
        match = SAFE_EXCLUSION.fullmatch(value)
        if match is None:
            raise ValueError("unsafe_reproducibility_metadata")
        counts = {name: int(token) for name, token in match.groupdict().items()}
        if counts["used"] + counts["missing"] != counts["input"]:
            raise ValueError("unsafe_reproducibility_metadata")
        cleaned.append(value)
    return "; ".join(dict.fromkeys(cleaned)) if cleaned else fallback


def _safe_transformations(values: Sequence[str], fallback: str) -> str:
    if any(not isinstance(value, str) or value not in SAFE_TRANSFORMATIONS for value in values):
        raise ValueError("unsafe_reproducibility_metadata")
    cleaned = list(dict.fromkeys(values))
    return "; ".join(cleaned) if cleaned else fallback


def _safe_software_versions(versions: dict[str, str]) -> str:
    if set(versions) != SAFE_SOFTWARE_COMPONENTS:
        raise ValueError("unsafe_reproducibility_metadata")
    if any(
        not isinstance(value, str) or SAFE_VERSION.fullmatch(value) is None
        for value in versions.values()
    ):
        raise ValueError("unsafe_reproducibility_metadata")
    return "; ".join(f"{name} {versions[name]}" for name in sorted(versions))


def _add_reproducibility_appendix(
    document: DocumentType,
    bundle: AnalysisBundle,
    language: Language,
) -> None:
    labels = RESULT_LABELS[language]
    provenance = bundle.provenance
    if not SAFE_FINGERPRINT.fullmatch(provenance.data_fingerprint):
        raise ValueError("invalid_report_fingerprint")
    if type(provenance.plan_version) is not int or provenance.plan_version < 1:
        raise ValueError("unsafe_reproducibility_metadata")
    versions = _safe_software_versions(bundle.reproducibility)
    exclusions = _safe_exclusions(provenance.exclusions, labels["not_applicable"])
    transformations = _safe_transformations(
        provenance.transformations, labels["not_applicable"]
    )
    document.add_heading(labels["reproducibility"], level=2)
    entries = (
        labels["plan_version"].format(value=provenance.plan_version),
        labels["fingerprint"].format(
            value=provenance.data_fingerprint[:FINGERPRINT_PREFIX_LENGTH]
        ),
        labels["software"].format(value=versions),
        labels["exclusions"].format(value=exclusions),
        labels["transformations"].format(value=transformations),
        labels["random_seed"],
    )
    for entry in entries:
        paragraph = document.add_paragraph(entry)
        paragraph.paragraph_format.keep_together = True


def _validate_inputs(plan: AnalysisPlan, bundle: AnalysisBundle) -> None:
    if plan.blocking_errors:
        raise ValueError("report_plan_has_blocking_errors")
    if bundle.provenance.plan_version != plan.version:
        raise ValueError("report_result_plan_mismatch")
    expected = [item.id for item in plan.items]
    if len(expected) != len(set(expected)) or set(expected) != set(bundle.results):
        raise ValueError("report_result_plan_mismatch")
    for item in plan.items:
        result = bundle.results[item.id]
        if result.method != item.method or result.provenance.plan_version != plan.version:
            raise ValueError("report_result_plan_mismatch")
        if result.provenance.data_fingerprint != bundle.provenance.data_fingerprint:
            raise ValueError("report_result_provenance_mismatch")


def build_results_docx(
    project: LocalProject,
    brief: StudyBrief,
    plan: AnalysisPlan,
    bundle: AnalysisBundle,
    figures: Sequence[FigureArtifact],
    language: Language,
    destination: Path,
) -> Path:
    """Build one localized, deterministic, manuscript-ready Results DOCX."""
    if language not in RESULT_LABELS:
        raise ValueError("unsupported_report_language")
    if not isinstance(project, LocalProject):
        raise TypeError("project_must_be_local_project")
    if not isinstance(brief, StudyBrief):
        raise TypeError("brief_must_be_study_brief")
    _validate_inputs(plan, bundle)
    destination = Path(destination)
    if destination.suffix.lower() != ".docx":
        raise ValueError("report_destination_must_be_docx")
    if destination.is_symlink():
        raise ValueError("unsafe_report_destination")
    destination.parent.mkdir(parents=True, exist_ok=True)

    labels = RESULT_LABELS[language]
    document = Document()
    _configure_page(document)
    _configure_styles(document)
    _configure_metadata(document)
    document.add_heading(labels["heading"], level=1)
    _add_results_narrative(document, brief, plan, bundle, language)
    _add_result_table(document, plan, bundle, language)
    _add_figures(document, figures, language)
    _add_warnings(document, plan, bundle, language)
    _add_reproducibility_appendix(document, bundle, language)
    document.save(destination)
    return destination


__all__ = ["RESULT_LABELS", "build_results_docx", "format_p_value"]

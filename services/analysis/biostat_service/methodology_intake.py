"""Metodoloji dokümanından deterministik metin çıkarma."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document

from biostat_service.data_intake import sha256_file


MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_DOCUMENT_CHARS = 200_000

SUPPORTED_FORMATS = frozenset({"docx", "pdf", "txt", "md"})


class MethodologyIntakeError(ValueError):
    """Stable, value-free failure code for document intake."""


@dataclass(frozen=True)
class MethodologyDocument:
    source_sha256: str
    source_format: str
    text: str
    char_count: int
    truncated: bool
    warnings: tuple[str, ...]


def _source_format(path: Path) -> str:
    suffix = path.suffix.lower().lstrip(".")
    if suffix not in SUPPORTED_FORMATS:
        raise MethodologyIntakeError("unsupported_format")
    return suffix


def _read_docx(path: Path) -> str:
    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _read_plain(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_document(path: Path) -> MethodologyDocument:
    source_format = _source_format(path)

    try:
        size = path.stat().st_size
    except OSError:
        raise MethodologyIntakeError("unreadable_document") from None
    if size > MAX_DOCUMENT_BYTES:
        raise MethodologyIntakeError("file_too_large")

    try:
        if source_format == "docx":
            raw = _read_docx(path)
        elif source_format == "pdf":
            raw = _read_pdf(path)
        else:
            raw = _read_plain(path)
    except MethodologyIntakeError:
        raise
    except Exception:  # noqa: BLE001 - any reader failure is one stable code
        raise MethodologyIntakeError("unreadable_document") from None

    text = raw.strip()
    if not text:
        raise MethodologyIntakeError("no_extractable_text")

    warnings: list[str] = []
    truncated = len(text) > MAX_DOCUMENT_CHARS
    if truncated:
        text = text[:MAX_DOCUMENT_CHARS]
        warnings.append("document_truncated")

    return MethodologyDocument(
        source_sha256=sha256_file(path),
        source_format=source_format,
        text=text,
        char_count=len(text),
        truncated=truncated,
        warnings=tuple(warnings),
    )
